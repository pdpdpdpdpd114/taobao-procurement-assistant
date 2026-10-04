from __future__ import annotations

import argparse
import shutil
from copy import deepcopy
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "src" / "assets" / "purchase-template.docx"


def set_cell_text(cell, text: str, *, align=None, size=10.5):
    if align is None:
        align = WD_ALIGN_PARAGRAPH.CENTER
    cell.text = text
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for paragraph in cell.paragraphs:
        paragraph.alignment = align
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1
        for run in paragraph.runs:
            run.font.size = Pt(size)
            run.font.name = "宋体"
            run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")


def prevent_row_split(row, repeat_header=False):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)
    if repeat_header:
        header = OxmlElement("w:tblHeader")
        header.set(qn("w:val"), "true")
        tr_pr.append(header)


def set_row_height(row, points: float):
    row.height = Pt(points)
    row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST


def set_complete_grid(table):
    """Remove inherited border suppressions and write a complete grid.

    The original form contains `top:nil` / `bottom:nil` overrides on the
    department and group-leader value cells.  These explicit overrides win over
    the table style and make the horizontal line disappear in Word.  Applying
    the same border definition to every physical cell keeps the complete form
    grid stable across Word and WPS.
    """
    for tr in table._tbl.tr_lst:
        for tc in tr.tc_lst:
            tc_pr = tc.get_or_add_tcPr()
            existing = tc_pr.find(qn("w:tcBorders"))
            if existing is not None:
                tc_pr.remove(existing)
            borders = OxmlElement("w:tcBorders")
            for edge in ("top", "left", "bottom", "right"):
                border = OxmlElement(f"w:{edge}")
                border.set(qn("w:val"), "single")
                border.set(qn("w:sz"), "4")
                border.set(qn("w:space"), "0")
                border.set(qn("w:color"), "000000")
                borders.append(border)
            tc_pr.append(borders)


def delete_row(table, index: int):
    row = table.rows[index]
    table._tbl.remove(row._tr)


def add_appendix(document: Document, source_table):
    start = document.add_paragraph()
    start.paragraph_format.space_after = Pt(0)
    start.add_run("{#hasAppendix}")

    heading = document.add_paragraph()
    heading.paragraph_format.page_break_before = True
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading.paragraph_format.space_after = Pt(8)
    run = heading.add_run("采购明细附表")
    run.bold = True
    run.font.size = Pt(16)
    run.font.name = "宋体"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")

    meta = document.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    meta.paragraph_format.space_after = Pt(6)
    meta_run = meta.add_run("部门：{department}    申请人：{applicant}    采购日期：{purchaseDate}    附表小计：{appendixSubtotal}")
    meta_run.font.size = Pt(10.5)

    appendix = document.add_table(rows=0, cols=6)
    appendix._tbl.remove(appendix._tbl.tblPr)
    appendix._tbl.insert(0, deepcopy(source_table._tbl.tblPr))
    appendix._tbl.append(deepcopy(source_table.rows[8]._tr))
    appendix._tbl.append(deepcopy(source_table.rows[9]._tr))

    headers = ["物品", "型号", "数量", "单价", "运费", "总价"]
    for idx, text in enumerate(headers):
        set_cell_text(appendix.rows[0].cells[idx], text)
    prevent_row_split(appendix.rows[0], repeat_header=True)

    values = [
        "{#appendixItems}{name}",
        "{model}",
        "{quantity}",
        "{unitPrice}",
        "{freight}",
        "{total}{/appendixItems}",
    ]
    for idx, text in enumerate(values):
        set_cell_text(
            appendix.rows[1].cells[idx],
            text,
            align=WD_ALIGN_PARAGRAPH.LEFT if idx < 2 else WD_ALIGN_PARAGRAPH.CENTER,
            size=9.5 if idx < 2 else 10.5,
        )
    prevent_row_split(appendix.rows[1])
    set_row_height(appendix.rows[1], 24)

    end = document.add_paragraph()
    end.paragraph_format.space_before = Pt(0)
    end.paragraph_format.space_after = Pt(0)
    end.add_run("{/hasAppendix}")


def main():
    parser = argparse.ArgumentParser(description="默认复用已提交的模板；--source 从原始六列采购表重新生成模板。")
    parser.add_argument("--source", type=Path, help="原始采购表 .docx 的路径（可选）")
    args = parser.parse_args()
    if args.source is None:
        if not OUTPUT.is_file():
            raise FileNotFoundError("未找到已提交的模板；请恢复 src/assets/purchase-template.docx，或通过 --source 指定原始采购表。")
        print(f"复用已提交模板：{OUTPUT}")
        return

    reference = args.source.expanduser().resolve()
    if not reference.is_file():
        raise FileNotFoundError(f"未找到原始采购表：{reference}")
    if reference == OUTPUT.resolve():
        raise ValueError("--source 必须指向原始采购表，不能指向已生成的模板。")

    # Loading an existing template needs only the Python standard library.
    global Document, WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
    global WD_ALIGN_PARAGRAPH, OxmlElement, qn, Pt
    try:
        from docx import Document
        from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Pt
    except ModuleNotFoundError as error:
        raise SystemExit("从原始表格重建模板需要 python-docx；请用所选 Python 运行 -m pip install python-docx==1.2.0。") from error

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(reference, OUTPUT)
    document = Document(OUTPUT)
    table = document.tables[0]
    table.autofit = False
    set_complete_grid(table)

    # Keep the fill date at the end of the completed application, rather than
    # between the title and the main table.
    fill_date = document.paragraphs[1]
    fill_date.text = "填报日期：{fillDate}"
    fill_date.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fill_date.paragraph_format.space_before = Pt(0)
    fill_date.paragraph_format.space_after = Pt(4)
    for run in fill_date.runs:
        run.font.size = Pt(10.5)
        run.font.name = "宋体"
        run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")

    field_slots = {
        (0, 1): "{department}",
        (0, 3): "{applicant}",
        (0, 5): "{purchaser}",
        (1, 1): "{leader}",
        (1, 3): "{captain}",
        (1, 5): "{teacher}",
        (2, 1): "{amount}",
        (2, 3): "{purchaseDate}",
        (2, 5): "{arrivalDate}",
        (3, 1): "{inventory}",
        (3, 3): "{custodian}",
        (3, 5): "{invoiceDate}",
        (4, 1): "{annualBudget}",
        (4, 3): "{currentSpent}",
        (4, 5): "{remainingBudget}",
    }
    # The original form uses manual spaces in several labels. Rebuild every
    # fixed label so Word renders a complete, stable six-column table.
    labels = {
        (0, 0): "部门", (0, 2): "申请人", (0, 4): "采购人",
        (1, 0): "组长", (1, 2): "队长/\n副队长", (1, 4): "总负责老师",
        (2, 0): "金额", (2, 2): "采购时间", (2, 4): "到货时间",
        (3, 0): "库存情况", (3, 2): "使用/保管\n人", (3, 4): "发票上交时间\n(采购人)",
        (4, 0): "年度预算", (4, 2): "目前金额", (4, 4): "剩余金额",
        (5, 0): "用途", (6, 0): "制作成品",
    }
    for (row_index, cell_index), text in labels.items():
        set_cell_text(table.rows[row_index].cells[cell_index], text)
    for (row_index, cell_index), tag in field_slots.items():
        set_cell_text(table.rows[row_index].cells[cell_index], tag)

    set_cell_text(table.rows[7].cells[0], "采购内容（不够另附清单）")
    for cell_index, text in enumerate(["物品", "型号", "数量", "单价", "运费", "总价"]):
        set_cell_text(table.rows[8].cells[cell_index], text)
    set_cell_text(table.rows[5].cells[1], "{purpose}", align=WD_ALIGN_PARAGRAPH.LEFT)
    set_cell_text(table.rows[6].cells[1], "{finishedProduct}", align=WD_ALIGN_PARAGRAPH.LEFT)
    set_row_height(table.rows[5], 34)
    set_row_height(table.rows[6], 34)

    # Keep one repeatable item row; Docxtemplater clones it for mainItems.
    for index in range(len(table.rows) - 1, 9, -1):
        delete_row(table, index)
    main_values = [
        "{#mainItems}{name}",
        "{model}",
        "{quantity}",
        "{unitPrice}",
        "{freight}",
        "{total}{/mainItems}",
    ]
    for idx, text in enumerate(main_values):
        set_cell_text(
            table.rows[9].cells[idx],
            text,
            align=WD_ALIGN_PARAGRAPH.LEFT if idx < 2 else WD_ALIGN_PARAGRAPH.CENTER,
            size=9.5 if idx < 2 else 10.5,
        )
    set_row_height(table.rows[9], 24)

    for row_index, row in enumerate(table.rows):
        prevent_row_split(row, repeat_header=row_index == 8)

    add_appendix(document, table)
    document._body._element.append(fill_date._p)

    document.core_properties.title = "车队经费使用申请模板"
    document.core_properties.subject = "淘宝购物车采购申请表"
    document.core_properties.author = "车队采购工具"
    document.core_properties.last_modified_by = "车队采购工具"
    document.save(OUTPUT)
    print(f"已生成模板：{OUTPUT}")


if __name__ == "__main__":
    main()
